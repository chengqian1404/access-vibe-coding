import io
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.utils.logger import logger
from app.utils.helpers import format_duration, get_timestamp_str


class ReportGenerator:
    """Generates Markdown and Excel reports from analysis results."""

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def get_report_path(self, recording_id: int) -> str:
        from app.config.settings import settings
        return str(Path(settings.ANALYSES_PATH) / f"recording_{recording_id}")

    def generate_markdown(self, analysis: Dict, recording: Dict) -> str:
        """Build a full Markdown report string."""
        lines: List[str] = []

        title = recording.get("title", "直播录制")
        created_at = recording.get("created_at", "")
        duration = recording.get("duration")
        duration_str = format_duration(duration) if duration else "未知"

        lines += [
            f"# 直播学习分析报告",
            f"",
            f"**录制标题：** {title}",
            f"**分析时间：** {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC",
            f"**录制时长：** {duration_str}",
            f"",
            "---",
            "",
            "## 一、内容概述",
            "",
            analysis.get("summary") or "（暂无摘要）",
            "",
        ]

        key_topics = analysis.get("key_topics") or []
        if key_topics:
            lines += ["## 二、核心话题", ""]
            for i, topic in enumerate(key_topics, 1):
                lines.append(f"{i}. {topic}")
            lines.append("")

        speech_style = analysis.get("speech_style") or {}
        if speech_style:
            lines += [
                "## 三、讲话风格分析",
                "",
                f"- **语气风格：** {speech_style.get('tone', '—')}",
                f"- **语速：** {speech_style.get('pace', '—')}",
                f"- **用词水平：** {speech_style.get('vocabulary_level', '—')}",
                f"- **专业度评分：** {speech_style.get('professionalism_score', '—')} / 10",
                "",
            ]
            techniques = speech_style.get("persuasion_techniques") or []
            if techniques:
                lines += ["**营销话术：**", ""]
                for t in techniques:
                    lines.append(f"- {t}")
                lines.append("")

            strengths = speech_style.get("strengths") or []
            if strengths:
                lines += ["**优点：**", ""]
                for s in strengths:
                    lines.append(f"- {s}")
                lines.append("")

            improvements = speech_style.get("improvement_suggestions") or speech_style.get("improvements") or []
            if improvements:
                lines += ["**改进建议：**", ""]
                for imp in improvements:
                    lines.append(f"- {imp}")
                lines.append("")

        qa_pairs = analysis.get("qa_pairs") or []
        lines += [
            "## 四、问答汇总",
            "",
            f"共 **{len(qa_pairs)}** 对问答",
            "",
        ]
        for i, qa in enumerate(qa_pairs, 1):
            ts = qa.get("timestamp")
            ts_str = f"（{format_duration(ts)}）" if ts else ""
            category = qa.get("category", "other")
            lines += [
                f"### Q{i}{ts_str} [{category}]",
                f"",
                f"**问：** {qa.get('question', '')}",
                f"",
                f"**答：** {qa.get('answer') or '（主播未直接回答）'}",
                "",
            ]

        key_points = analysis.get("key_points") or []
        if key_points:
            lines += ["## 五、关键信息点", ""]
            for kp in sorted(key_points, key=lambda k: -(k.get("importance") or 3)):
                importance = kp.get("importance", 3)
                stars = "★" * importance + "☆" * (5 - importance)
                ts = kp.get("timestamp")
                ts_str = f" [{format_duration(ts)}]" if ts else ""
                lines.append(f"- {stars}{ts_str} **{kp.get('category', '')}** {kp.get('content', '')}")
            lines.append("")

        audience_eng = analysis.get("audience_engagement") or ""
        if audience_eng:
            lines += ["## 六、观众互动分析", "", audience_eng, ""]

        lines += [
            "---",
            "",
            f"*报告生成时间：{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC*",
        ]

        return "\n".join(lines)

    def generate_excel(self, analysis: Dict, recording: Dict) -> bytes:
        """Generate a multi-sheet Excel workbook and return as bytes."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl is required for Excel generation")

        wb = openpyxl.Workbook()

        # ---- Sheet 1: Summary ----
        ws_summary = wb.active
        ws_summary.title = "概述"
        title_font = Font(bold=True, size=14)
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="4472C4")

        ws_summary.append(["直播学习分析报告"])
        ws_summary["A1"].font = title_font
        ws_summary.append([])
        ws_summary.append(["录制标题", recording.get("title", "")])
        ws_summary.append(["录制时长", format_duration(recording.get("duration") or 0)])
        ws_summary.append(["分析时间", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])
        ws_summary.append(["问答总数", analysis.get("total_qa_count", len(analysis.get("qa_pairs") or []))])
        ws_summary.append([])
        ws_summary.append(["内容摘要", analysis.get("summary", "")])
        ws_summary.column_dimensions["A"].width = 16
        ws_summary.column_dimensions["B"].width = 60

        # ---- Sheet 2: Q&A ----
        ws_qa = wb.create_sheet("问答列表")
        qa_headers = ["序号", "时间", "分类", "问题", "回答", "置信度"]
        ws_qa.append(qa_headers)
        for cell in ws_qa[1]:
            cell.font = header_font
        for i, qa in enumerate(analysis.get("qa_pairs") or [], 1):
            ts = qa.get("timestamp")
            ws_qa.append([
                i,
                format_duration(ts) if ts else "",
                qa.get("category", ""),
                qa.get("question", ""),
                qa.get("answer") or "",
                f"{(qa.get('confidence') or 0):.0%}",
            ])
        for col, width in zip("ABCDEF", (6, 10, 14, 40, 50, 10)):
            ws_qa.column_dimensions[col].width = width

        # ---- Sheet 3: Key Points ----
        ws_kp = wb.create_sheet("关键信息")
        kp_headers = ["序号", "重要度", "分类", "内容", "时间"]
        ws_kp.append(kp_headers)
        for cell in ws_kp[1]:
            cell.font = header_font
        for i, kp in enumerate(analysis.get("key_points") or [], 1):
            ts = kp.get("timestamp")
            ws_kp.append([
                i,
                kp.get("importance", 3),
                kp.get("category", ""),
                kp.get("content", ""),
                format_duration(ts) if ts else "",
            ])
        for col, width in zip("ABCDE", (6, 10, 14, 60, 10)):
            ws_kp.column_dimensions[col].width = width

        # ---- Sheet 4: Speech Style ----
        ws_style = wb.create_sheet("讲话风格")
        speech = analysis.get("speech_style") or {}
        style_rows = [
            ("语气风格", speech.get("tone", "")),
            ("语速", speech.get("pace", "")),
            ("用词水平", speech.get("vocabulary_level", "")),
            ("专业度评分", speech.get("professionalism_score", "")),
            ("营销话术", "\n".join(speech.get("persuasion_techniques") or [])),
            ("号召行动", "\n".join(speech.get("call_to_action") or [])),
            ("优点", "\n".join(speech.get("strengths") or [])),
            ("改进建议", "\n".join(speech.get("improvement_suggestions") or speech.get("improvements") or [])),
        ]
        for row in style_rows:
            ws_style.append(row)
        ws_style.column_dimensions["A"].width = 16
        ws_style.column_dimensions["B"].width = 60

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def save_markdown(self, analysis: Dict, recording: Dict) -> str:
        """Save the markdown report to disk and return the file path."""
        report_dir = Path(self.get_report_path(recording.get("id", 0)))
        report_dir.mkdir(parents=True, exist_ok=True)
        ts = get_timestamp_str()
        path = report_dir / f"report_{ts}.md"
        content = self.generate_markdown(analysis, recording)
        path.write_text(content, encoding="utf-8")
        logger.info("Markdown report saved: %s", path)
        return str(path)

    def save_excel(self, analysis: Dict, recording: Dict) -> str:
        """Save the Excel report to disk and return the file path."""
        report_dir = Path(self.get_report_path(recording.get("id", 0)))
        report_dir.mkdir(parents=True, exist_ok=True)
        ts = get_timestamp_str()
        path = report_dir / f"report_{ts}.xlsx"
        data = self.generate_excel(analysis, recording)
        path.write_bytes(data)
        logger.info("Excel report saved: %s", path)
        return str(path)
